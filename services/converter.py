#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Call Center - Universal File Converter Engine
Supports:
  - PDF -> DOCX
  - DOCX -> PDF
  - TXT -> PDF, TXT -> DOCX
  - Images (JPG, PNG, WEBP) -> PDF
  - Images format conversion (JPG <-> PNG <-> WEBP)
"""

import sys
import os
import argparse

def convert_pdf_to_docx(src, dst):
    from pdf2docx import Converter
    cv = Converter(src)
    cv.convert(dst)
    cv.close()

def setup_unicode_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    search_dirs = [
        "/opt/call-center/fonts",
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "fonts"),
        "C:\\Windows\\Fonts",
        "/usr/share/fonts/truetype/dejavu",
        "/usr/share/fonts/dejavu",
        "/usr/share/fonts/truetype/liberation"
    ]

    def find_font(filename):
        for d in search_dirs:
            full = os.path.join(d, filename)
            if os.path.exists(full):
                return full
        return None

    # 1. Times New Roman
    times_reg = find_font("times.ttf")
    times_bd = find_font("timesbd.ttf")
    times_it = find_font("timesi.ttf")
    times_bi = find_font("timesbi.ttf")

    if times_reg and times_bd:
        try:
            pdfmetrics.registerFont(TTFont("TimesNewRoman", times_reg))
            pdfmetrics.registerFont(TTFont("TimesNewRoman-Bold", times_bd))
            pdfmetrics.registerFont(TTFont("TimesNewRoman-Italic", times_it if times_it else times_reg))
            pdfmetrics.registerFont(TTFont("TimesNewRoman-BoldItalic", times_bi if times_bi else times_bd))
            pdfmetrics.registerFontFamily(
                "TimesNewRoman",
                normal="TimesNewRoman",
                bold="TimesNewRoman-Bold",
                italic="TimesNewRoman-Italic",
                boldItalic="TimesNewRoman-BoldItalic"
            )
        except Exception as e:
            sys.stderr.write(f"TimesNewRoman register error: {e}\n")

    # 2. Arial
    arial_reg = find_font("arial.ttf")
    arial_bd = find_font("arialbd.ttf")
    if arial_reg and arial_bd:
        try:
            pdfmetrics.registerFont(TTFont("Arial", arial_reg))
            pdfmetrics.registerFont(TTFont("Arial-Bold", arial_bd))
            pdfmetrics.registerFontFamily("Arial", normal="Arial", bold="Arial-Bold")
        except Exception as e:
            sys.stderr.write(f"Arial register error: {e}\n")

    # 3. Fallback DejaVu / Liberation
    dejavu_reg = find_font("DejaVuSans.ttf") or find_font("LiberationSans-Regular.ttf")
    dejavu_bd = find_font("DejaVuSans-Bold.ttf") or find_font("LiberationSans-Bold.ttf")
    if dejavu_reg:
        try:
            pdfmetrics.registerFont(TTFont("UniFont", dejavu_reg))
            pdfmetrics.registerFont(TTFont("UniFont-Bold", dejavu_bd if dejavu_bd else dejavu_reg))
            pdfmetrics.registerFontFamily("UniFont", normal="UniFont", bold="UniFont-Bold")
        except Exception as e:
            sys.stderr.write(f"UniFont register error: {e}\n")

    if "TimesNewRoman" in pdfmetrics.getRegisteredFontNames():
        return "TimesNewRoman", "TimesNewRoman-Bold"
    elif "Arial" in pdfmetrics.getRegisteredFontNames():
        return "Arial", "Arial-Bold"
    elif "UniFont" in pdfmetrics.getRegisteredFontNames():
        return "UniFont", "UniFont-Bold"

    return "Helvetica", "Helvetica-Bold"

def convert_docx_to_pdf(src, dst):
    import html
    import io
    import tempfile
    import docx
    from PIL import Image as PILImage
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    font_reg, font_bold = setup_unicode_fonts()
    PAGE_W, PAGE_H = A4
    content_width = PAGE_W - 72  # 523 pt

    doc = docx.Document(src)

    with tempfile.TemporaryDirectory(prefix="cc_docx_imgs_") as img_dir:
        # 1. Map related image parts
        rel_parts = {}
        for r_id, part in doc.part.related_parts.items():
            if hasattr(part, 'blob') and hasattr(part, 'content_type') and part.content_type.startswith('image/'):
                img_path = os.path.join(img_dir, f"{r_id}.png")
                try:
                    pil_img = PILImage.open(io.BytesIO(part.blob))
                    if pil_img.mode in ('RGBA', 'LA', 'P'):
                        bg = PILImage.new('RGB', pil_img.size, (255, 255, 255))
                        bg.paste(pil_img, mask=pil_img.split()[-1] if pil_img.mode == 'RGBA' else None)
                        bg.save(img_path, format='PNG')
                    elif pil_img.mode != 'RGB':
                        pil_img.convert('RGB').save(img_path, format='PNG')
                    else:
                        pil_img.save(img_path, format='PNG')
                    rel_parts[r_id] = {
                        'path': img_path,
                        'orig_w': pil_img.size[0],
                        'orig_h': pil_img.size[1]
                    }
                except Exception as img_err:
                    sys.stderr.write(f"Image decode warning {r_id}: {img_err}\n")

        # 2. Extract absolute floating images (e.g. badge/profile photo in top corner)
        absolute_images = []
        for pic in doc.element.body.xpath('.//*[local-name()="imagedata"]'):
            r_id = pic.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            parent = pic.getparent()
            style = parent.attrib.get('style', '') if parent is not None else ''
            if r_id and r_id in rel_parts:
                props = {}
                for item in style.split(';'):
                    if ':' in item:
                        k, v = item.split(':', 1)
                        props[k.strip().lower()] = v.strip()
                if 'absolute' in props.get('position', ''):
                    w_pt = 86.0
                    h_pt = 103.0
                    if 'width' in props:
                        try: w_pt = float(props['width'].replace('pt', '').replace('px', ''))
                        except: pass
                    if 'height' in props:
                        try: h_pt = float(props['height'].replace('pt', '').replace('px', ''))
                        except: pass

                    x_pt = PAGE_W - 36 - w_pt - 10
                    y_top = 20.0
                    if 'margin-top' in props:
                        try: y_top = float(props['margin-top'].replace('pt', ''))
                        except: pass

                    absolute_images.append({
                        'path': rel_parts[r_id]['path'],
                        'x': x_pt,
                        'y': PAGE_H - 36 - y_top - h_pt,
                        'w': w_pt,
                        'h': h_pt
                    })

        def on_first_page_cb(canvas, doc_template):
            for ai in absolute_images:
                try:
                    canvas.drawImage(ai['path'], ai['x'], ai['y'], width=ai['w'], height=ai['h'], preserveAspectRatio=True)
                except Exception as draw_err:
                    sys.stderr.write(f"Failed to draw image: {draw_err}\n")

        # Styles
        styles = getSampleStyleSheet()
        normal = ParagraphStyle(
            'Normal_Custom',
            fontName=font_reg,
            fontSize=9.5,
            leading=13.5,
            textColor=colors.HexColor('#0f172a')
        )
        title_style = ParagraphStyle(
            'Title_Custom',
            fontName=font_bold,
            fontSize=15,
            leading=19,
            alignment=1,
            spaceAfter=6
        )
        sub_title_style = ParagraphStyle(
            'SubTitle_Custom',
            fontName=font_bold,
            fontSize=11.5,
            leading=15.5,
            spaceBefore=6,
            spaceAfter=4
        )

        pdf = SimpleDocTemplate(
            dst,
            pagesize=A4,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        story = []

        # Iterate elements in document body order
        for el in doc.element.body:
            tag = el.tag.split('}')[-1]
            if tag == 'p':
                p = docx.text.paragraph.Paragraph(el, doc)
                raw_text = p.text.strip()

                # Inline images
                inline_imgs = []
                for b in el.xpath('.//*[local-name()="blip"]'):
                    r_id = b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if r_id and r_id in rel_parts:
                        inline_imgs.append(rel_parts[r_id])
                for pic in el.xpath('.//*[local-name()="imagedata"]'):
                    r_id = pic.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                    parent = pic.getparent()
                    style = parent.attrib.get('style', '') if parent is not None else ''
                    if 'absolute' not in style and r_id and r_id in rel_parts:
                        inline_imgs.append(rel_parts[r_id])

                for im in inline_imgs:
                    max_w = content_width
                    disp_w = min(im['orig_w'], max_w)
                    ratio = disp_w / float(im['orig_w'])
                    disp_h = im['orig_h'] * ratio
                    if disp_h > 350:
                        disp_h = 350
                        disp_w = im['orig_w'] * (disp_h / float(im['orig_h']))
                    story.append(RLImage(im['path'], width=disp_w, height=disp_h))
                    story.append(Spacer(1, 4))

                if not raw_text:
                    story.append(Spacer(1, 3))
                    continue

                clean_text = html.escape(raw_text)
                clean_lower = clean_text.lower()
                style_name = (p.style.name or '').lower()

                if 'heading 1' in style_name or 'title' in style_name or 'маълумотнома' in clean_lower or 'резюме' in clean_lower:
                    story.append(Paragraph(clean_text, title_style))
                elif 'heading 2' in style_name or 'меҳнат фаолияти' in clean_lower or 'маълумот' in clean_lower or 'яқин' in clean_lower:
                    story.append(Paragraph(clean_text, sub_title_style))
                elif '\t' in raw_text:
                    cols = [c.strip() for c in raw_text.split('\t') if c.strip()]
                    if len(cols) == 2:
                        # Keep max width bounded to leave space for right profile photo
                        col_w1 = 175
                        col_w2 = 235
                        t_row = Table(
                            [[Paragraph(html.escape(cols[0]), normal), Paragraph(html.escape(cols[1]), normal)]],
                            colWidths=[col_w1, col_w2]
                        )
                        t_row.setStyle(TableStyle([
                            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                            ('LEFTPADDING', (0, 0), (-1, -1), 0),
                            ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                            ('TOPPADDING', (0, 0), (-1, -1), 1),
                            ('BOTTOMPADDING', (0, 0), (-1, -1), 1),
                        ]))
                        story.append(t_row)
                    elif len(cols) > 2:
                        cw = content_width / len(cols)
                        t_row = Table(
                            [[Paragraph(html.escape(c), normal) for c in cols]],
                            colWidths=[cw] * len(cols)
                        )
                        t_row.setStyle(TableStyle([
                            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                            ('LEFTPADDING', (0, 0), (-1, -1), 0),
                            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                            ('TOPPADDING', (0, 0), (-1, -1), 1),
                            ('BOTTOMPADDING', (0, 0), (-1, -1), 1),
                        ]))
                        story.append(t_row)
                    else:
                        story.append(Paragraph(clean_text, normal))
                        story.append(Spacer(1, 2))
                else:
                    story.append(Paragraph(clean_text, normal))
                    story.append(Spacer(1, 2))

            elif tag == 'tbl':
                table = docx.table.Table(el, doc)
                table_data = []
                for row in table.rows:
                    row_cells = []
                    for cell in row.cells:
                        ctext = cell.text.strip()
                        cell_imgs = []
                        for b in cell._tc.xpath('.//*[local-name()="blip"]'):
                            r_id = b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                            if r_id and r_id in rel_parts:
                                cell_imgs.append(rel_parts[r_id])
                        for pic in cell._tc.xpath('.//*[local-name()="imagedata"]'):
                            r_id = pic.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                            if r_id and r_id in rel_parts:
                                cell_imgs.append(rel_parts[r_id])

                        flowables = []
                        for im in cell_imgs:
                            flowables.append(RLImage(im['path'], width=60, height=75))
                        if ctext:
                            flowables.append(Paragraph(html.escape(ctext), normal))
                        if not flowables:
                            flowables.append(Paragraph("", normal))
                        row_cells.append(flowables if len(flowables) > 1 else flowables[0])
                    table_data.append(row_cells)

                if table_data:
                    num_cols = len(table.columns) if table.columns else len(table_data[0])
                    col_widths = None
                    if num_cols == 5:
                        col_widths = [
                            content_width * 0.14,
                            content_width * 0.23,
                            content_width * 0.21,
                            content_width * 0.19,
                            content_width * 0.23
                        ]
                    report_table = Table(table_data, colWidths=col_widths)
                    report_table.setStyle(TableStyle([
                        ('FONTNAME', (0, 0), (-1, -1), font_reg),
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f1f5f9')),
                        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#94a3b8')),
                        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                        ('TOPPADDING', (0, 0), (-1, -1), 3),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
                    ]))
                    story.append(Spacer(1, 6))
                    story.append(report_table)
                    story.append(Spacer(1, 6))

        if not story:
            story.append(Paragraph("(Bo'sh hujjat)", normal))

        pdf.build(story, onFirstPage=on_first_page_cb)

def convert_txt_to_pdf(src, dst):
    import html
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    font_reg, font_bold = setup_unicode_fonts()

    pdf = SimpleDocTemplate(dst, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    normal = styles['Normal']
    normal.fontName = font_reg
    normal.fontSize = 10
    normal.leading = 14

    story = []
    with open(src, 'r', encoding='utf-8', errors='ignore') as f:
        for line in f:
            line_str = line.strip()
            if not line_str:
                story.append(Spacer(1, 6))
            else:
                story.append(Paragraph(html.escape(line_str), normal))
                story.append(Spacer(1, 2))
    
    if not story:
        story.append(Paragraph("(Bo'sh matn)", normal))

    pdf.build(story)

def convert_txt_to_docx(src, dst):
    import docx
    doc = docx.Document()
    with open(src, 'r', encoding='utf-8', errors='ignore') as f:
        for line in f:
            doc.add_paragraph(line.rstrip('\r\n'))
    doc.save(dst)

def convert_img_to_pdf(src, dst):
    from PIL import Image
    img = Image.open(src)
    if img.mode in ('RGBA', 'LA', 'P'):
        bg = Image.new('RGB', img.size, (255, 255, 255))
        bg.paste(img, mask=img.split()[-1] if img.mode == 'RGBA' else None)
        img = bg
    elif img.mode != 'RGB':
        img = img.convert('RGB')
    img.save(dst, 'PDF', resolution=100.0)

def convert_img_to_img(src, dst, target_format):
    from PIL import Image
    img = Image.open(src)
    fmt = target_format.upper()
    if fmt in ('JPG', 'JPEG'):
        fmt = 'JPEG'
        if img.mode in ('RGBA', 'LA', 'P'):
            bg = Image.new('RGB', img.size, (255, 255, 255))
            bg.paste(img, mask=img.split()[-1] if img.mode == 'RGBA' else None)
            img = bg
        elif img.mode != 'RGB':
            img = img.convert('RGB')
        img.save(dst, format=fmt, quality=92)
    elif fmt == 'PNG':
        img.save(dst, format='PNG', optimize=True)
    elif fmt == 'WEBP':
        img.save(dst, format='WEBP', quality=90)
    else:
        img.save(dst)

def main():
    parser = argparse.ArgumentParser(description="Call Center File Converter")
    parser.add_argument('input', help="Path to input file")
    parser.add_argument('output', help="Path to output file")
    parser.add_argument('--target', required=True, help="Target format extension (pdf, docx, png, jpg, etc.)")
    args = parser.parse_args()

    src = os.path.abspath(args.input)
    dst = os.path.abspath(args.output)
    target = args.target.lower().replace('.', '')

    if not os.path.exists(src):
        print(f"ERROR: Input file does not exist: {src}", file=sys.stderr)
        sys.exit(1)

    src_ext = os.path.splitext(src)[1].lower().replace('.', '')

    # Dispatch conversion
    try:
        if src_ext == 'pdf' and target in ('docx', 'doc'):
            convert_pdf_to_docx(src, dst)
        elif src_ext in ('docx', 'doc') and target == 'pdf':
            convert_docx_to_pdf(src, dst)
        elif src_ext == 'txt' and target == 'pdf':
            convert_txt_to_pdf(src, dst)
        elif src_ext == 'txt' and target in ('docx', 'doc'):
            convert_txt_to_docx(src, dst)
        elif src_ext in ('jpg', 'jpeg', 'png', 'webp', 'bmp') and target == 'pdf':
            convert_img_to_pdf(src, dst)
        elif src_ext in ('jpg', 'jpeg', 'png', 'webp', 'bmp') and target in ('jpg', 'jpeg', 'png', 'webp'):
            convert_img_to_img(src, dst, target)
        else:
            print(f"ERROR: Unsupported conversion: {src_ext} -> {target}", file=sys.stderr)
            sys.exit(2)

        if not os.path.exists(dst) or os.path.getsize(dst) == 0:
            print("ERROR: Output file was not generated or is empty.", file=sys.stderr)
            sys.exit(3)

        print(f"SUCCESS: Converted {src_ext} to {target}, size={os.path.getsize(dst)}")
        sys.exit(0)
    except Exception as e:
        print(f"ERROR: Conversion failed: {str(e)}", file=sys.stderr)
        sys.exit(4)

if __name__ == '__main__':
    main()
